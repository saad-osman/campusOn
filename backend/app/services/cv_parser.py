"""Feature 1: CV/transcript text extraction and profile-field extraction.

Files are parsed from memory and never written to disk (Section 8: "store
extracted text only; delete the original after parsing"). With an API key the
text goes to EXTRACTION_MODEL; otherwise a rule-based extractor runs. Either
way the student reviews the suggested values before anything is saved.
"""
import io
import re
from datetime import date

from app.services.llm import LLMUnavailable, complete_json, llm_enabled, load_prompt

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class CVParseError(Exception):
    pass


def extract_cv_text(filename: str, data: bytes) -> str:
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise CVParseError("Upload a PDF or DOCX file")
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif ext == ".docx":
            import docx

            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        else:
            text = data.decode("utf-8", errors="replace")
    except CVParseError:
        raise
    except Exception as e:
        raise CVParseError("Couldn't read that file. Is it a valid PDF or DOCX?") from e
    text = re.sub(r"[ \t]+", " ", text).strip()
    if len(text) < 30:
        raise CVParseError("No readable text found. Scanned PDFs need OCR first; try a DOCX or text-based PDF.")
    return text


KNOWN_SKILLS = [
    "Python", "Java", "C++", "C", "JavaScript", "TypeScript", "React", "Node.js", "SQL", "R", "MATLAB",
    "PyTorch", "TensorFlow", "Keras", "scikit-learn", "Pandas", "NumPy", "OpenCV", "ROS", "Docker",
    "Kubernetes", "AWS", "Git", "Linux", "Rust", "Go", "Excel", "Tableau", "Figma", "Verilog", "LaTeX",
    "Hugging Face", "Spark", "FastAPI", "Django", "Flutter",
]
KNOWN_INTERESTS = [
    "machine learning", "deep learning", "computer vision", "natural language processing", "robotics",
    "reinforcement learning", "data science", "public policy", "cybersecurity", "distributed systems",
    "human-computer interaction", "bioinformatics", "renewable energy", "quantum computing", "economics",
    "sustainability", "healthcare", "climate", "finance", "blockchain", "embedded systems",
]
KNOWN_MAJORS = [
    "Computer Science", "Electrical and Electronics Engineering", "Electronics and Communication",
    "Electrical Engineering", "Mechanical Engineering", "Chemical Engineering", "Civil Engineering",
    "Biotechnology", "Mathematics", "Physics", "Chemistry", "Economics", "Data Science",
    "Information Technology", "Biology", "Business", "Public Policy",
]
UAE_HINTS = ["uae", "united arab emirates", "dubai", "abu dhabi", "sharjah", "ajman"]
NATIONALITIES = ["Indian", "Pakistani", "Emirati", "Egyptian", "Jordanian", "Bangladeshi", "Sri Lankan",
                 "Filipino", "Nepali", "British", "American", "Canadian", "Syrian", "Lebanese", "Sudanese"]


def extract_profile_rule_based(text: str, today: date | None = None) -> dict:
    today = today or date.today()
    lower = text.lower()
    out: dict = {}

    if re.search(r"\b(ph\.?d|doctoral)\b", lower) and not re.search(r"\b(pursue|apply|applying) (a )?ph\.?d", lower):
        out["degree_level"] = "phd"
    elif re.search(r"\b(m\.?sc|m\.?tech|master'?s|m\.?e\.)\b", lower):
        out["degree_level"] = "masters"
    elif re.search(r"\b(b\.?e\.?|b\.?tech|b\.?sc|bachelor'?s?|undergraduate)\b", lower):
        out["degree_level"] = "bachelors"

    m = re.search(r"\b([1-5])(?:st|nd|rd|th)[\s-]*year\b", lower) or re.search(r"\byear\s*([1-5])\b", lower)
    if m:
        out["year_of_study"] = int(m.group(1))
    else:
        g = re.search(r"(?:expected|graduat\w*|class of)[^\d]{0,25}(20\d\d)", lower)
        if g and out.get("degree_level", "bachelors") == "bachelors":
            years_left = int(g.group(1)) - today.year + (1 if today.month >= 7 else 0)
            out["year_of_study"] = max(1, min(4, 4 - years_left + 1))

    for major in KNOWN_MAJORS:
        if major.lower() in lower:
            out["major"] = major
            break

    m = re.search(r"\b(?:c?gpa)\b[^\d]{0,15}(\d{1,2}(?:\.\d{1,2})?)\s*(?:/|out of)\s*(\d{1,2}(?:\.\d)?)", lower)
    if m:
        out["cgpa"], out["cgpa_scale"] = float(m.group(1)), float(m.group(2))
    else:
        m = re.search(r"\b(?:c?gpa)\b[^\d]{0,15}(\d{1,2}\.\d{1,2})", lower)
        if m:
            value = float(m.group(1))
            out["cgpa"], out["cgpa_scale"] = value, 4.0 if value <= 4.0 else 10.0

    m = re.search(r"nationality[:\s]+([A-Za-z ]{3,30})", text, re.I)
    if m:
        out["nationality"] = m.group(1).strip().split("\n")[0].title()
    else:
        for nat in NATIONALITIES:
            if re.search(rf"\b{nat}\b", text):
                out["nationality"] = nat
                break
    if any(h in lower for h in UAE_HINTS):
        out["country_of_residence"] = "UAE"

    tests = {}
    m = re.search(r"ielts[^\d]{0,15}(\d(?:\.\d)?)", lower)
    if m:
        tests["IELTS"] = float(m.group(1))
    m = re.search(r"toefl[^\d]{0,15}(\d{2,3})", lower)
    if m:
        tests["TOEFL"] = float(m.group(1))
    if tests:
        out["english_tests"] = tests

    skills = [s for s in KNOWN_SKILLS if re.search(rf"(?<![\w+#.]){re.escape(s)}(?![\w+#])", text)]
    if skills:
        out["skills"] = skills[:15]

    interests = []
    m = re.search(r"(?:research )?interests?[ \t]*:[ \t]*([^\n]+)", text, re.I)
    if m:
        interests = [i.strip(" .;") for i in re.split(r",|;|/", m.group(1)) if 2 < len(i.strip()) < 50]
    for kw in KNOWN_INTERESTS:
        if kw in lower and kw not in [i.lower() for i in interests]:
            interests.append(kw)
    if interests:
        out["interests"] = interests[:10]
    return out


PROFILE_SYSTEM = load_prompt("profile_extraction.md")

PROFILE_KEYS = ["degree_level", "year_of_study", "major", "cgpa", "cgpa_scale", "nationality",
                "country_of_residence", "english_tests", "skills", "interests"]


def _clean(data: dict) -> dict:
    out = {}
    for key in PROFILE_KEYS:
        value = data.get(key)
        if value in (None, "", [], {}):
            continue
        if key == "degree_level" and value not in ("bachelors", "masters", "phd"):
            continue
        if key in ("year_of_study",):
            try:
                value = int(value)
            except (TypeError, ValueError):
                continue
        if key in ("cgpa", "cgpa_scale"):
            try:
                value = float(value)
            except (TypeError, ValueError):
                continue
        if key == "english_tests":
            value = {str(k).upper(): float(v) for k, v in value.items() if isinstance(v, (int, float))}
        if key in ("skills", "interests"):
            value = [str(v) for v in value if v][:15]
        out[key] = value
    return out


def extract_profile(cv_text: str) -> tuple[dict, str]:
    if llm_enabled():
        try:
            return _clean(complete_json(PROFILE_SYSTEM, f"CV text:\n{cv_text[:12000]}", max_tokens=1200)), "llm"
        except LLMUnavailable:
            pass
    return _clean(extract_profile_rule_based(cv_text)), "rule_based"
